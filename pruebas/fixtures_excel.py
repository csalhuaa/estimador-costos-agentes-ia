"""Fresh synthetic inputs, never copied from or fitted to the original case."""

from datetime import date
from pathlib import Path

from openpyxl import Workbook

from modelo.constantes import NOMBRES_EM, constantes_predeterminadas

SHEETS = (
    "LEEME",
    "proyecto",
    "parametros_ia",
    "estudios",
    "presupuesto_inicial",
    "costo_real_sprint",
    "constantes",
    "verificacion",
    "rangos",
)
PARAMETERS = (
    "r",
    "rho_rev",
    "rho_retr",
    "c_api_h",
    "delta",
    "h_def",
    "p_lic",
    "h_cap",
    "h_int",
)
RANGES = (
    (0.1, 0.3, 0.5),
    (0.05, 0.1, 0.2),
    (0.02, 0.05, 0.1),
    (0.1, 0.2, 0.3),
    (0.01, 0.02, 0.04),
    (1, 2, 3),
    (10, 20, 30),
    (4, 8, 16),
    (2, 4, 8),
)


def libro_sintetico(*, caso=True):
    """Match source rows/symbols with hardcoded, explicitly declared assumptions."""
    w = Workbook()
    w.active.title = SHEETS[0]
    for name in SHEETS[1:]:
        w.create_sheet(name)
    w["LEEME"]["A1"] = "SYNTHETIC ONLY — no primary-study evidence"
    w["LEEME"]["A2"] = (
        "Edit project C5:C34 and agent D6:I14; example values are assumptions."
    )
    c = constantes_predeterminadas()
    symbols = [
        "nombre",
        "moneda",
        "Size",
        *c.sf_nominal,
        *NOMBRES_EM,
        "SCED_EM",
        "SCED_pct",
        "c_pm",
        "f",
        "T",
        "n_u",
    ]
    values = [
        "SYNTHETIC demonstration",
        "USD",
        32,
        *c.sf_nominal.values(),
        *([1] * 16),
        1,
        100,
        1000,
        "0,59",
        None,
        10,
    ]
    for row, (symbol, value) in enumerate(zip(symbols, values), 5):
        w["proyecto"].cell(row, 2, symbol)
        w["proyecto"].cell(row, 3, value)
    for row, (symbol, numbers) in enumerate(zip(PARAMETERS, RANGES), 6):
        s = w["parametros_ia"]
        s.cell(row, 1, symbol)
        for column, number in enumerate(numbers, 4):
            s.cell(row, column, number)
        s.cell(row, 7, "Supuesto declarado")
        s.cell(row, 9, "Synthetic illustrative assumption; not empirical evidence")
    for row in range(6, 26):
        w["estudios"].cell(row, 1, f"E{row - 5:02}")
    s = w["presupuesto_inicial"]
    s["B4"], s["B8"] = "USD", "No"
    for row in range(6, 18):
        w["costo_real_sprint"].cell(row, 1, row - 5)
    if caso:
        s["B5"], s["B6"], s["B7"] = date(2026, 1, 1), 32, "No"
        s["B9"] = "Synthetic historical cash assumption"
        for col, value in enumerate(
            (
                "Synthetic budget",
                "Programming",
                1000,
                120000,
                "Declared",
                "Synthetic note",
                "Sí",
            ),
            1,
        ):
            s.cell(13, col, value)
        s = w["costo_real_sprint"]
        for col, value in enumerate(
            (
                1,
                date(2026, 1, 2),
                "2026-01-16",
                19800,
                None,
                None,
                None,
                None,
                None,
                None,
                3000,
                500,
                200,
                100,
                "Synthetic observation",
                "Si",
            ),
            1,
        ):
            if value is not None:
                s.cell(6, col, value)
    names = [
        "A",
        "B",
        "C",
        "D",
        "horas_por_PM",
        *(f"{x} nominal" for x in c.sf_nominal),
        "EM nominal",
        "SCED% nominal",
    ]
    values = [c.A, c.B, c.C, c.D, c.horas_por_pm, *c.sf_nominal.values(), 1, 100]
    for row, (name, value) in enumerate(zip(names, values), 5):
        w["constantes"].cell(row, 1, name)
        w["constantes"].cell(row, 2, value)
    return w


def guardar_sintetico(path: Path, *, caso=True):
    """Create a new input only; existing files are never overwritten."""
    w = libro_sintetico(caso=caso)
    try:
        with path.open("xb") as stream:
            w.save(stream)
    finally:
        w.close()
    return path


def resultados_sinteticos(path: Path):
    """Test/demo assembly only; no public pipeline or duplicate model arithmetic."""
    from modelo.carga import cargar_entradas
    from modelo.cocomo import estimar_base, verificar_cocomo
    from modelo.contraste import contrastar
    from modelo.costo_real import calcular_costo_real
    from modelo.escenarios import correr_escenarios, decidir, sensibilidad
    from modelo.parametros_ia import contrastar_r_con_estudios
    from modelo.tipos import Decision, Resultados
    from dataclasses import replace

    e = cargar_entradas(path)
    e = replace(
        e, avisos=e.avisos + contrastar_r_con_estudios(e.parametros, e.estudios)
    )
    b = estimar_base(e.proyecto, e.constantes)
    scenarios = correr_escenarios(e.proyecto, e.constantes, e.parametros)
    d = Decision(
        scenarios,
        sensibilidad(e.proyecto, e.constantes, e.parametros),
        20,
        decidir(scenarios),
    )
    cr = calcular_costo_real(e.caso, b.t_h) if e.caso else None
    contrast = (
        contrastar(
            b.C_trad,
            scenarios[1].resultado.C_ia,
            cr,
            scenarios,
            verificar_cocomo(e.constantes),
        )
        if cr
        else None
    )
    return Resultados(e, b, d, cr, contrast, date(2026, 10, 24))
