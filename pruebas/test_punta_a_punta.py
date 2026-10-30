"""End-to-end verification against Section 13 acceptance criteria."""

from pathlib import Path

from openpyxl import load_workbook
import pytest

from modelo.cocomo import verificar_cocomo
from modelo.constantes import constantes_predeterminadas
from modelo.errores import ErrorDeEntrada
from modelo.pipeline import ejecutar
from modelo.reporte import escribir_reporte
from pruebas.fixtures_excel import guardar_sintetico, libro_sintetico


@pytest.fixture
def golden_path(tmp_path):
    return guardar_sintetico(tmp_path / "ejemplo_dorado.xlsx", caso=True)


@pytest.fixture
def no_case_path(tmp_path):
    return guardar_sintetico(tmp_path / "ejemplo_sin_caso.xlsx", caso=False)


def test_criterio_verificar_cocomo_nominal():
    """Criterion: verificar_cocomo evaluates to True with nominal manual constants."""
    assert verificar_cocomo(constantes_predeterminadas()) is True


def test_criterio_excel_con_errores_muestra_todos_a_la_vez(tmp_path):
    """Criterion: An Excel file with multiple missing inputs reports all problems simultaneously."""
    template_path = tmp_path / "invalid.xlsx"
    workbook = libro_sintetico(caso=False)
    for row in workbook["proyecto"].iter_rows(
        min_row=5, max_row=34, min_col=3, max_col=3
    ):
        row[0].value = None
    for row in workbook["parametros_ia"].iter_rows(
        min_row=6, max_row=14, min_col=4, max_col=9
    ):
        for cell in row:
            cell.value = None
    workbook.save(template_path)
    workbook.close()
    with pytest.raises(ErrorDeEntrada) as exc_info:
        ejecutar(template_path)

    problemas = exc_info.value.problemas
    assert len(problemas) >= 40
    assert any(p.codigo == "E01" and p.hoja == "proyecto" for p in problemas)
    assert any(p.codigo == "E01" and p.hoja == "parametros_ia" for p in problemas)
    assert any(p.codigo == "E05" and p.hoja == "parametros_ia" for p in problemas)
    assert all(p.codigo and p.hoja and p.mensaje for p in problemas)


def test_criterio_ejecucion_con_y_sin_caso(golden_path, no_case_path):
    """Criterion: Without case, M1-M5/M8 execute without error; with case, M6 and M7 are included."""
    res_sin_caso = ejecutar(no_case_path)
    assert res_sin_caso.base is not None
    assert res_sin_caso.decision is not None
    assert res_sin_caso.costo_real is None
    assert res_sin_caso.contraste is None

    res_con_caso = ejecutar(golden_path)
    assert res_con_caso.base is not None
    assert res_con_caso.decision is not None
    assert res_con_caso.costo_real is not None
    assert res_con_caso.contraste is not None
    assert res_con_caso.contraste.verificacion_cocomo_ok is True


def test_criterio_trazabilidad_formulas_y_fuentes(golden_path, tmp_path):
    """Criterion: Every reported figure includes its formula (F...) and its source cell coordinate."""
    res = ejecutar(golden_path)
    out_dir = tmp_path / "reportes_e2e"
    reportes = escribir_reporte(res, out_dir)
    xlsx_path = [p for p in reportes if p.name == "reporte.xlsx"][0]

    wb = load_workbook(xlsx_path, data_only=True)
    assert "trazabilidad" in wb.sheetnames
    ws_trace = wb["trazabilidad"]
    headers = [ws_trace.cell(1, col).value for col in range(1, 5)]
    assert headers == ["cifra", "fórmula", "entradas usadas", "valor"]

    rows = list(ws_trace.iter_rows(min_row=2, values_only=True))
    assert len(rows) >= 50
    for r in rows:
        cifra, formula, entradas, valor = r
        assert cifra and str(cifra).strip()
        assert formula and str(formula).strip()
        assert entradas and str(entradas).strip()
        assert valor is not None or "None" in str(formula)
    wb.close()


def test_criterio_supuestos_declarados_visibles(golden_path, tmp_path):
    """Criterion: Declared assumptions appear in the report artifacts and markdown output."""
    res = ejecutar(golden_path)
    assert (
        len(res.entradas.supuestos) == 9
    )  # all 9 agent parameters declared as assumptions in fixture

    out_dir = tmp_path / "reportes_supuestos"
    reportes = escribir_reporte(res, out_dir)
    md_content = [p for p in reportes if p.name == "reporte.md"][0].read_text(
        encoding="utf-8"
    )
    assert "Supuestos declarados" in md_content
    for sup in res.entradas.supuestos:
        assert sup in md_content


def test_criterio_coherencia_f34_y_exactitud_25_porciento(golden_path):
    """Criterion: Accuracy follows F34 and stability transports M5 agreement."""
    res = ejecutar(golden_path)
    assert res.contraste is not None
    # F34: err_estimacion_ia <= 0.25
    assert res.contraste.cumple_exactitud == (res.contraste.err_estimacion_ia <= 0.25)
    # Decision stability flag evaluated
    assert isinstance(res.contraste.decision_estable, bool)
    assert res.contraste.decision_estable is res.decision.robusta


def test_criterio_ausencia_algoritmos_aleatorios_o_bootstrap():
    """Criterion: Codebase contains zero non-deterministic random, bootstrap, or permutation code."""
    root_modelo = Path(__file__).resolve().parents[1] / "modelo"
    forbidden_terms = [
        "bootstrap",
        "permutation",
        "random.choice",
        "random.sample",
        "np.random",
    ]
    for py_file in root_modelo.glob("*.py"):
        content = py_file.read_text(encoding="utf-8").lower()
        for term in forbidden_terms:
            assert term not in content, (
                f"Found forbidden term '{term}' in {py_file.name}"
            )
