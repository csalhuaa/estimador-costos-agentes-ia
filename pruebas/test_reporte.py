"""M8 snapshots agree across formats and never execute spreadsheet strings."""

from dataclasses import replace
from hashlib import sha256
import json

from openpyxl import load_workbook
from PIL import Image
import pytest

from modelo.reporte import a_diccionario, escribir_reporte
from pruebas.fixtures_excel import (
    guardar_sintetico,
    libro_sintetico,
    resultados_sinteticos,
)


@pytest.fixture
def result(tmp_path):
    return resultados_sinteticos(guardar_sintetico(tmp_path / "synthetic.xlsx"))


@pytest.mark.parametrize("sheet", ["proyecto", "parametros_ia", "constantes"])
def test_report_provenance_survives_inserted_row(tmp_path, sheet):
    workbook = libro_sintetico(caso=False)
    workbook[sheet].insert_rows(6)
    path = tmp_path / "grown.xlsx"
    workbook.save(path)
    workbook.close()
    before = sha256(path.read_bytes()).digest()
    try:
        data = a_diccionario(resultados_sinteticos(path))
        cells = {
            (c["hoja"], c["celda"]): c["valor"]
            for c in data["entradas"]["fuente"]["celdas"]
        }
        if sheet == "proyecto":
            assert cells["proyecto", "C35"] == 10
            assert "proyecto!C35" in data["escenarios"][1]["costos"]["C_lic"]["fuente"]
        elif sheet == "parametros_ia":
            assert cells["parametros_ia", "E15"] == 4
            assert (
                data["escenarios"][1]["valores"]["h_int"]["fuente"]
                == "parametros_ia!D15:I15 (h_int)"
            )
        else:
            assert cells["constantes", "B17"] == 100
            assert (
                "constantes!B17"
                in data["contraste"]["verificacion_cocomo_ok"]["fuente"]
                if data["contraste"]
                else not data["entradas"]["fuente"]["constantes_predeterminadas"]
            )
    finally:
        assert sha256(path.read_bytes()).digest() == before


def test_four_cross_format_snapshots_and_valid_png(result, tmp_path):
    source = tmp_path / "synthetic.xlsx"
    before = sha256(source.read_bytes()).hexdigest()
    paths = escribir_reporte(result, tmp_path / "reports")
    assert [p.name for p in paths] == [
        "reporte.xlsx",
        "reporte.md",
        "resultados.json",
        "sensibilidad.png",
    ]
    assert set(paths[0].parent.iterdir()) == set(paths)
    data = json.loads(paths[2].read_text(encoding="utf-8"))
    assert data == a_diccionario(result)
    figure = data["base"]["C_trad"]
    assert figure["valor"] == result.base.C_trad
    assert figure["formula"] == "F7: C_trad = PM * c_pm"
    assert "proyecto!C31" in figure["fuente"]
    assert data["resumen"]["veredicto"]["valor"] == result.decision.veredicto
    assert data["resumen"]["fecha_ejecucion"] == "2026-10-24"
    assert data["resumen"]["archivo_entrada"] == "synthetic.xlsx"
    assert data["costo_real"]["H_real_4"]["valor"] is None
    assert (
        data["resumen"]["mayor_efecto"]["valor"]
        == result.decision.sensibilidad[0].parametro
    )
    markdown = paths[1].read_text(encoding="utf-8")
    assert "132,910.95" in markdown
    assert "F7: C_trad = PM * c_pm" in markdown
    assert "Supuestos" in markdown and "Avisos" in markdown
    assert "sensible a los parámetros" in markdown
    summary = data["resumen"]
    keys = list(summary)
    start = keys.index("veredicto")
    assert keys[start : start + 3] == ["veredicto", "robusta", "etiqueta_robustez"]
    assert summary["robusta"]["valor"] is result.decision.robusta
    assert summary["etiqueta_robustez"]["valor"] == "sensible a los parámetros"
    for name in keys[start : start + 3]:
        assert "F30" in summary[name]["formula"]
        assert summary[name]["fuente"] == summary["veredicto"]["fuente"]
    md_rows = [line for line in markdown.splitlines() if line.startswith("| resumen.")]
    assert [line.split(" | ")[0] for line in md_rows[:3]] == [
        "| resumen.veredicto",
        "| resumen.robusta",
        "| resumen.etiqueta_robustez",
    ]
    w = load_workbook(paths[0], data_only=False)
    assert w.sheetnames == [
        "resumen",
        "escenarios",
        "sensibilidad",
        "contraste",
        "supuestos",
        "trazabilidad",
    ]
    assert [c.value for c in w["trazabilidad"][1]] == [
        "cifra",
        "fórmula",
        "entradas usadas",
        "valor",
    ]
    rows = list(w["trazabilidad"].values)
    summary_rows = [r for r in w["resumen"].values if str(r[0]).startswith("resumen.")]
    assert [r[0] for r in summary_rows[:3]] == [
        "resumen.veredicto",
        "resumen.robusta",
        "resumen.etiqueta_robustez",
    ]
    assert summary_rows[1][3] is result.decision.robusta
    assert any(
        r[0] == "base.C_trad" and r[3] == pytest.approx(result.base.C_trad, rel=1e-15)
        for r in rows
    )
    for s in w:
        for row in s:
            for c in row:
                assert c.data_type != "f"
                assert c.font.name == "Arial"
    w.close()
    with Image.open(paths[3]) as image:
        assert image.format == "PNG"
        assert image.width >= 800 and image.height >= 400
        image.verify()
    assert sha256(source.read_bytes()).hexdigest() == before
    assert "ruta_interna" not in paths[2].read_text(encoding="utf-8")
    assert str(tmp_path) not in paths[2].read_text(encoding="utf-8")


def test_no_case_omits_contrast_and_strict_json_preserves_none(tmp_path):
    result = resultados_sinteticos(
        guardar_sintetico(tmp_path / "no-case.xlsx", caso=False)
    )
    paths = escribir_reporte(result, tmp_path / "reports")
    w = load_workbook(paths[0], read_only=True)
    assert "contraste" not in w.sheetnames
    w.close()
    data = json.loads(paths[2].read_text(encoding="utf-8"))
    assert data["contraste"] is None and data["costo_real"] is None


@pytest.mark.parametrize("prefix", ["=", "+", "-", "@"])
def test_untrusted_strings_are_literal_cells(result, tmp_path, prefix):
    text = prefix + 'HYPERLINK("https://invalid.example")'
    e = replace(
        result.entradas,
        proyecto=replace(result.entradas.proyecto, nombre=text),
        supuestos=(text,),
    )
    result = replace(result, entradas=e)
    paths = escribir_reporte(result, tmp_path / "reports")
    w = load_workbook(paths[0], data_only=False)
    matches = [c for s in w for row in s for c in row if c.value == text]
    assert matches
    assert all(c.data_type == "s" for c in matches)
    assert not any(c.data_type == "f" for s in w for row in s for c in row)
    w.close()


@pytest.mark.parametrize(
    "name", ["reporte.xlsx", "reporte.md", "resultados.json", "sensibilidad.png"]
)
def test_any_existing_target_refuses_before_any_writes(result, tmp_path, name):
    out = tmp_path / "reports"
    out.mkdir()
    existing = out / name
    existing.write_bytes(b"user bytes")
    with pytest.raises(FileExistsError):
        escribir_reporte(result, out)
    assert list(out.iterdir()) == [existing]
    assert existing.read_bytes() == b"user bytes"


def test_source_collision_is_refused(result, tmp_path):
    path = tmp_path / "resultados.json"
    path.write_bytes(b"source sentinel")
    source = replace(result.entradas.fuente, ruta_interna=str(path.resolve()))
    result = replace(result, entradas=replace(result.entradas, fuente=source))
    with pytest.raises(FileExistsError):
        escribir_reporte(result, tmp_path)
    assert path.read_bytes() == b"source sentinel"
    assert not (tmp_path / "reporte.xlsx").exists()


def test_failure_cleans_only_newly_owned_outputs(result, tmp_path, monkeypatch):
    from pathlib import Path

    original = Path.open
    out = tmp_path / "reports"
    out.mkdir()
    sentinel = out / "keep.txt"
    sentinel.write_bytes(b"keep")

    def fail(path, mode="r", *args, **kwargs):
        if path.name == "resultados.json" and mode == "xb":
            raise OSError("Synthetic disk failure")
        return original(path, mode, *args, **kwargs)

    monkeypatch.setattr(Path, "open", fail)
    with pytest.raises(OSError, match="Synthetic disk failure"):
        escribir_reporte(result, out)
    assert list(out.iterdir()) == [sentinel]
    assert sentinel.read_bytes() == b"keep"


def test_nonfinite_result_is_rejected_before_writing(result, tmp_path):
    result = replace(result, base=replace(result.base, C_trad=float("nan")))
    with pytest.raises(ValueError):
        escribir_reporte(result, tmp_path / "reports")
    assert not (tmp_path / "reports").exists()


def test_legacy_constructors_report_explicit_missing_provenance(result, tmp_path):
    result = replace(result, entradas=replace(result.entradas, fuente=None))
    d = a_diccionario(result)
    assert "no disponible" in d["base"]["C_trad"]["fuente"].lower()
    paths = escribir_reporte(result, tmp_path / "reports")
    assert len(paths) == 4


def test_default_constants_and_runtime_arguments_never_get_fake_coordinates(result):
    source = replace(result.entradas.fuente, constantes_predeterminadas=("A",))
    result = replace(result, entradas=replace(result.entradas, fuente=source))
    d = a_diccionario(result)
    assert "predeterminado declarado: A" in d["base"]["PM"]["fuente"]
    assert "runtime: x_pct" in d["sensibilidad"][0]["efecto_bajo"]["fuente"]
    assert "constantes!B5" not in d["base"]["PM"]["fuente"]


def test_zero_cost_none_and_raw_precision_are_preserved(result):
    from modelo.escenarios import correr_escenarios, decidir
    from modelo.tipos import Decision, ParametrosIA, Trio

    params = ParametrosIA(
        *(
            Trio(v, v, v, "Supuesto declarado", (), "Synthetic")
            for v in (0.3, 0, 0, 0, 0, 0, 0, 0, 0)
        )
    )
    e = replace(result.entradas, parametros=params)
    scenarios = correr_escenarios(e.proyecto, e.constantes, params)
    result = replace(
        result, entradas=e, decision=Decision(scenarios, (), 20, *decidir(scenarios))
    )
    d = a_diccionario(result)
    assert d["escenarios"][1]["resultado"]["rbc"]["valor"] is None
    assert d["escenarios"][1]["resultado"]["recuperacion_meses"]["valor"] is None
    assert d["base"]["PM"]["valor"] == result.base.PM
    json.dumps(d, allow_nan=False)


def test_render_failure_leaves_no_outputs(result, tmp_path, monkeypatch):
    from matplotlib.backends.backend_agg import FigureCanvasAgg

    def fail(*args, **kwargs):
        raise OSError("Synthetic render failure")

    monkeypatch.setattr(FigureCanvasAgg, "print_png", fail)
    with pytest.raises(OSError, match="Synthetic render failure"):
        escribir_reporte(result, tmp_path / "reports")
    assert not (tmp_path / "reports").exists()


def test_robust_label_is_native_boolean_and_keeps_all_financial_values(result):
    before = a_diccionario(result)
    changed = replace(
        result, decision=replace(result.decision, veredicto="no conviene", robusta=True)
    )
    data = a_diccionario(changed)
    assert data["resumen"]["robusta"]["valor"] is True
    assert data["resumen"]["etiqueta_robustez"]["valor"] == "robusta"
    for key in (
        "base",
        "escenarios",
        "sensibilidad",
        "costo_real",
        "contraste",
        "entradas",
    ):
        assert data[key] == before[key]
