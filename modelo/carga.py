"""M0: read-only, aggregate validation of the nine-sheet input contract."""

from datetime import date, datetime
from hashlib import sha256
from math import isfinite
from pathlib import Path
from zipfile import BadZipFile
from xml.etree.ElementTree import ParseError

import defusedxml.ElementTree  # noqa: F401 -- required XML protection dependency
from defusedxml.common import DefusedXmlException
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from openpyxl.xml import DEFUSEDXML

from .constantes import NOMBRES_EM, constantes_predeterminadas
from .errores import ErrorDeEntrada, Problema
from .tipos import (
    Aviso,
    CasoValidacion,
    CeldaFuente,
    Constantes,
    Entradas,
    Estudio,
    FuenteEntrada,
    ParametrosIA,
    Proyecto,
    Rubro,
    Sprint,
    Trio,
)

HOJAS = (
    "LEEME",
    "proyecto",
    "parametros_ia",
    "estudios",
    "presupuesto_inicial",
    "costo_real_sprint",
    "constantes",
    "verificacion",
    "rangos",
)
PARAMETROS = (
    "r",
    "rho_rev",
    "rho_retr",
    "c_api_h",
    "delta",
    "h_def",
    "p_lic",
    "h_cap",
    "h_int",
)


class _Lectura:
    """Keep validation state local; no partial contracts escape on errors."""

    def __init__(self, workbook):
        self.w = workbook
        self.problemas = []
        self.avisos = []
        self.celdas = {}
        self.defaults = []

    def get(self, hoja, celda):
        value = self.w[hoja][celda].value if hoja in self.w.sheetnames else None
        self.celdas[hoja, celda] = CeldaFuente(hoja, celda, value)
        return value

    def error(self, code, sheet, cell, message, value):
        self.problemas.append(
            Problema(code, sheet, cell, f"{message} (valor: {value}).")
        )

    def warn(self, code, sheet, cell, message):
        self.avisos.append(Aviso(code, sheet, cell, message))

    def text(self, sheet, cell, *, required=True):
        value = self.get(sheet, cell)
        if value is None or str(value).strip() == "":
            if required:
                self.error("E01", sheet, cell, "Campo obligatorio vacío", value)
            return ""
        return str(value).strip()

    def number(
        self,
        sheet,
        cell,
        *,
        optional=False,
        lower=0,
        strict=False,
        upper=None,
        integer=False,
    ):
        value = self.get(sheet, cell)
        if value is None or value == "":
            if not optional:
                self.error("E01", sheet, cell, "Campo obligatorio vacío", value)
            return None if optional else 0.0
        try:
            if isinstance(value, (bool, date)):
                raise ValueError
            number = float(str(value).strip().replace(",", "."))
            if not isfinite(number):
                raise ValueError
        except (ValueError, TypeError, OverflowError):
            self.error("E02", sheet, cell, "Se espera un número finito", value)
            return 0.0
        if (
            number < lower
            or (strict and number == lower)
            or (upper is not None and number > upper)
            or (integer and not number.is_integer())
        ):
            self.error("E03", sheet, cell, "Valor fuera de rango permitido", value)
        return number

    def yes(self, sheet, cell, *, review=False, optional=False):
        value = self.text(sheet, cell, required=not optional)
        normalized = {"si": "Si", "sí": "Si", "no": "No"}
        if review:
            normalized["por revisar"] = "Por revisar"
        if optional and not value:
            return ""
        if value.casefold() not in normalized:
            if value:
                self.error(
                    "E11",
                    sheet,
                    cell,
                    "Usar Si o No" + (" o Por revisar" if review else ""),
                    value,
                )
            return "No"
        return normalized[value.casefold()]

    def day(self, sheet, cell):
        value = self.get(sheet, cell)
        if value is None or value == "":
            return None
        try:
            if isinstance(value, datetime):
                return value.date()
            if isinstance(value, date):
                return value
            if isinstance(value, (int, float)) and not isinstance(value, bool):
                result = from_excel(value, self.w.epoch)
                if not isinstance(result, datetime):
                    raise ValueError
                return result.date()
            return date.fromisoformat(str(value))
        except (ValueError, TypeError, OverflowError):
            self.error("E13", sheet, cell, "Fecha inválida", value)
            return None

    def symbols(self, sheet, column, first=1, last=None):
        found = {}
        ws = self.w[sheet] if sheet in self.w.sheetnames else None
        if ws is None:
            return found
        max_r = ws.max_row if ws.max_row is not None else (first + 500)
        limit = max_r if last is None else max(last, max_r)
        consecutive_empty = 0
        for row in range(first, limit + 1):
            val = ws[f"{column}{row}"].value
            if val is not None and str(val).strip() != "":
                consecutive_empty = 0
                self.get(sheet, f"{column}{row}")
                symbol = str(val).strip()
                if symbol in found:
                    self.error("E07", sheet, f"{column}{row}", "Símbolo duplicado", val)
                found[symbol] = row
            else:
                consecutive_empty += 1
                if consecutive_empty >= 25 and (last is None or row >= last):
                    break
        return found

    def project(self):
        sheet = "proyecto"
        symbols = self.symbols(sheet, "B", 5)
        sf_names = tuple(constantes_predeterminadas().sf_nominal)
        expected = (
            "nombre",
            "moneda",
            "Size",
            *sf_names,
            *NOMBRES_EM,
            "SCED_EM",
            "SCED_pct",
            "c_pm",
            "f",
            "T",
            "n_u",
        )
        for name in expected:
            if name not in symbols:
                self.error("E07", sheet, "B5:B34", "Falta símbolo obligatorio", name)

        def cell(name):
            # Missing symbols have already been reported; no invented source cell.
            return f"C{symbols[name]}" if name in symbols else ""

        def num(name, **kwargs):
            return self.number(sheet, cell(name), **kwargs) if cell(name) else 0.0

        name = self.text(sheet, cell("nombre")) if cell("nombre") else ""
        currency = self.text(sheet, cell("moneda")) if cell("moneda") else ""
        if currency and currency != "USD":
            self.error("E06", sheet, cell("moneda"), "La moneda debe ser USD", currency)
        size = num("Size", strict=True)
        sf = {x: num(x, strict=True) for x in sf_names}
        em = {x: num(x, strict=True) for x in NOMBRES_EM}
        sced_em = num("SCED_EM", strict=True)
        sced_pct = num("SCED_pct", strict=True)
        c_pm = num("c_pm", strict=True)
        f = num("f", strict=True, upper=1)
        if f == 1:
            self.error("E03", sheet, cell("f"), "f debe ser menor que 1", f)
        if not 0.55 <= f <= 0.61:
            self.warn(
                "A01", sheet, cell("f"), "f fuera de 0,55 a 0,61; declarar supuesto"
            )
        horizon = num("T", optional=True, strict=True)
        users = num("n_u", lower=1, integer=True)
        return Proyecto(
            name,
            currency,
            size,
            sf,
            em,
            sced_em,
            sced_pct,
            c_pm,
            f,
            horizon,
            int(users),
        )

    def constants(self):
        d = constantes_predeterminadas()
        names = (
            "A",
            "B",
            "C",
            "D",
            "horas_por_PM",
            *(f"{x} nominal" for x in d.sf_nominal),
            "EM nominal",
            "SCED% nominal",
        )
        values = (
            d.A,
            d.B,
            d.C,
            d.D,
            d.horas_por_pm,
            *d.sf_nominal.values(),
            d.em_nominal,
            d.sced_pct_nominal,
        )
        symbols = self.symbols("constantes", "A", 5)
        result = {}
        for name, default in zip(names, values):
            row = symbols.get(name)
            cell = f"B{row}" if row else ""
            if not row or self.get("constantes", cell) in (None, ""):
                result[name] = default
                self.defaults.append(name)
                self.warn(
                    "A06",
                    "constantes",
                    cell,
                    f"{name}: valor predeterminado declarado de modelo/constantes.py",
                )
            else:
                result[name] = self.number("constantes", cell, strict=True)
        return Constantes(
            *(result[x] for x in names[:5]),
            {x: result[f"{x} nominal"] for x in d.sf_nominal},
            result["EM nominal"],
            result["SCED% nominal"],
        )

    def studies(self):
        studies = []
        for row in range(6, 26):
            values = [self.get("estudios", f"{col}{row}") for col in "ABCDEFGHIJKLMNO"]
            if not any(v not in (None, "") for v in values[1:]):
                continue
            study_id = self.text("estudios", f"A{row}")
            if study_id in {s.id for s in studies}:
                self.error(
                    "E05", "estudios", f"A{row}", "Id de estudio duplicado", study_id
                )
            studies.append(
                Estudio(
                    study_id,
                    str(values[1] or ""),
                    str(values[2] or ""),
                    str(values[3] or ""),
                    str(values[5] or ""),
                    self.number("estudios", f"L{row}", optional=True, upper=1),
                    self.yes("estudios", f"M{row}", review=True),
                )
            )
        return tuple(studies)

    def parameters(self, studies):
        symbols = self.symbols("parametros_ia", "A", 6)
        trios = []
        assumptions = []
        by_id = {s.id: s for s in studies}
        for name in PARAMETROS:
            if name not in symbols:
                self.error(
                    "E07", "parametros_ia", "A6:A14", "Falta símbolo obligatorio", name
                )
                trios.append(Trio(0, 0, 0, "", (), ""))
                continue
            row = symbols[name]
            numbers = [
                self.number(
                    "parametros_ia", f"{col}{row}", upper=1 if name == "r" else None
                )
                for col in "DEF"
            ]
            if not numbers[0] <= numbers[1] <= numbers[2]:
                self.error(
                    "E04",
                    "parametros_ia",
                    f"D{row}:F{row}",
                    "Se requiere mínimo ≤ esperado ≤ máximo",
                    numbers,
                )
            origin = self.text("parametros_ia", f"G{row}")
            ids = tuple(
                x.strip()
                for x in self.text("parametros_ia", f"H{row}", required=False).split(
                    ";"
                )
                if x.strip()
            )
            justification = self.text("parametros_ia", f"I{row}", required=False)
            if origin not in ("Estudio", "Supuesto declarado"):
                self.error(
                    "E05", "parametros_ia", f"G{row}", "Origen no reconocido", origin
                )
            elif origin == "Estudio" and not ids:
                self.error(
                    "E05", "parametros_ia", f"H{row}", "Estudio requiere ids", ids
                )
            elif origin == "Supuesto declarado":
                assumptions.append(name)
                if not justification:
                    self.error(
                        "E05",
                        "parametros_ia",
                        f"I{row}",
                        "Supuesto requiere justificación",
                        justification,
                    )
                self.warn(
                    "A04", "parametros_ia", f"G{row}", f"{name}: supuesto declarado"
                )
            for study_id in ids:
                if study_id not in by_id:
                    self.error(
                        "E05",
                        "parametros_ia",
                        f"H{row}",
                        "Id de estudio inexistente",
                        study_id,
                    )
                elif by_id[study_id].cumple != "Si":
                    self.warn(
                        "A08",
                        "parametros_ia",
                        f"H{row}",
                        f"{study_id}: criterios distintos de Si",
                    )
            trios.append(Trio(*numbers, origin, ids, justification))
        return ParametrosIA(*trios), tuple(assumptions)

    def case(self, project, constants):
        budget_sheet = "presupuesto_inicial"
        sprint_sheet = "costo_real_sprint"
        budget_rows = []
        sprint_rows = []
        # Retain supplemental metadata even when the optional case is absent.
        for row in range(4, 10):
            self.get(budget_sheet, f"B{row}")
        for row in range(13, 25):
            values = [self.get(budget_sheet, f"{col}{row}") for col in "ABCDEFG"]
            if any(v not in (None, "") for v in values):
                budget_rows.append(row)
        for row in range(6, 18):
            values = [self.get(sprint_sheet, f"{col}{row}") for col in "BCDEFGHKLMNOP"]
            if any(v not in (None, "") for v in values):
                sprint_rows.append(row)
        if not budget_rows and not sprint_rows:
            return None
        if bool(budget_rows) != bool(sprint_rows):
            self.error(
                "E08",
                budget_sheet,
                "A13:G24",
                "Caso incompleto: llenar ambas hojas",
                None,
            )
        currency = self.text(budget_sheet, "B4")
        if currency and currency != "USD":
            self.error("E06", budget_sheet, "B4", "La moneda debe ser USD", currency)
        self.day(budget_sheet, "B5")
        self.number(budget_sheet, "B6", optional=True, strict=True)
        self.yes(budget_sheet, "B7", optional=True)
        if self.yes(budget_sheet, "B8", optional=True) == "Si":
            self.warn("A07", budget_sheet, "B8", "El presupuesto consideró agentes")
        rubros = []
        for row in budget_rows:
            rubros.append(
                Rubro(
                    self.text(budget_sheet, f"A{row}"),
                    self.number(budget_sheet, f"C{row}", optional=True),
                    self.number(budget_sheet, f"D{row}"),
                    self.yes(budget_sheet, f"G{row}") == "Si",
                )
            )
        sprints = []
        seen = set()
        for row in sprint_rows:
            number = self.number(sprint_sheet, f"A{row}", lower=1, integer=True)
            if number in seen:
                self.error("E03", sprint_sheet, f"A{row}", "Sprint duplicado", number)
            seen.add(number)
            start = self.day(sprint_sheet, f"B{row}")
            end = self.day(sprint_sheet, f"C{row}")
            if start and end and end < start:
                self.error(
                    "E13", sprint_sheet, f"C{row}", "Fecha fin anterior a inicio", end
                )
            total = self.number(sprint_sheet, f"D{row}")
            activities = [
                self.number(sprint_sheet, f"{col}{row}", optional=True)
                for col in "EFGH"
            ]
            recorded = [x for x in activities if x is not None]
            hours4 = sum(recorded) if len(recorded) == 4 else None
            if sum(recorded) > total:
                self.error(
                    "E10",
                    sprint_sheet,
                    f"E{row}:H{row}",
                    "Horas de actividades mayores que total",
                    sum(recorded),
                )
            costs = [self.number(sprint_sheet, f"{col}{row}") for col in "KLMN"]
            included = self.yes(sprint_sheet, f"P{row}") == "Si"
            if included and hours4 is None:
                self.warn(
                    "A05",
                    sprint_sheet,
                    f"E{row}:H{row}",
                    "Horas por actividad incompletas; se omite F9",
                )
            sprints.append(
                Sprint(int(number), start, end, total, hours4, *costs, included)
            )
        if not any(r.en_alcance for r in rubros):
            self.error(
                "E09", budget_sheet, "G13:G24", "Se requiere rubro en alcance", None
            )
        if not any(s.hubo_programacion for s in sprints):
            self.error(
                "E09",
                sprint_sheet,
                "P6:P17",
                "Se requiere sprint con programación",
                None,
            )
        if rubros and sum(r.costo_usd for r in rubros if r.en_alcance) <= 0:
            self.error(
                "E03",
                budget_sheet,
                "D13:D24",
                "Presupuesto en alcance debe ser positivo",
                0,
            )
        if constants.horas_por_pm > 0 and sprints:
            total_real = sum(
                s.horas_totales * project.c_pm / constants.horas_por_pm
                + s.licencias
                + s.api
                + s.capacitacion
                + s.integracion
                for s in sprints
                if s.hubo_programacion
            )
            if total_real <= 0:
                self.error(
                    "E03",
                    sprint_sheet,
                    "D6:N17",
                    "Costo real en alcance debe ser positivo",
                    total_real,
                )
        return CasoValidacion(tuple(rubros), tuple(sprints))


def cargar_entradas(ruta: Path) -> Entradas:
    """Return a complete validated contract or all located E01–E13 problems."""
    ruta = Path(ruta)
    try:
        if ruta.suffix.lower() != ".xlsx":
            raise ValueError("Extensión no válida")
        if not DEFUSEDXML:
            raise ValueError("Protección XML no disponible")
        digest = sha256(ruta.read_bytes()).hexdigest()
        workbook = load_workbook(ruta, read_only=True, data_only=True, keep_links=False)
    except (
        OSError,
        ValueError,
        BadZipFile,
        KeyError,
        ParseError,
        DefusedXmlException,
    ) as exc:
        raise ErrorDeEntrada(
            [
                Problema(
                    "E12",
                    "archivo",
                    "",
                    f"No se puede abrir como .xlsx (valor: {ruta.name}).",
                )
            ]
        ) from exc
    try:
        reader = _Lectura(workbook)
        for sheet in HOJAS:
            if sheet not in workbook.sheetnames:
                reader.error("E07", sheet, "", "Falta hoja obligatoria", sheet)
        constants = reader.constants()
        project = reader.project()
        studies = reader.studies()
        parameters, assumptions = reader.parameters(studies)
        case = reader.case(project, constants)
        if reader.problemas:
            raise ErrorDeEntrada(reader.problemas)
        source = FuenteEntrada(
            ruta.name,
            digest,
            str(ruta.resolve()),
            tuple(reader.celdas.values()),
            tuple(reader.defaults),
        )
        return Entradas(
            constants,
            project,
            parameters,
            studies,
            case,
            tuple(reader.avisos),
            assumptions,
            source,
        )
    except (ParseError, DefusedXmlException, BadZipFile) as exc:
        raise ErrorDeEntrada(
            [
                Problema(
                    "E12", "archivo", "", f"XML de libro inválido (valor: {ruta.name})."
                )
            ]
        ) from exc
    finally:
        workbook.close()
