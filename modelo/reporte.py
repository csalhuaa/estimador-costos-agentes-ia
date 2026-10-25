"""M8: literal-safe static snapshots of existing M1–M7 results.

Expressions below are documentation, never executable calculations. All report
formats use the same annotated dictionary and retain unrounded model values.
"""

from dataclasses import fields, is_dataclass
from datetime import date
from io import BytesIO
import json
from pathlib import Path
import re

from matplotlib.backends.backend_agg import FigureCanvasAgg
from matplotlib.figure import Figure
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .tipos import Resultados

ARCHIVOS = ("reporte.xlsx", "reporte.md", "resultados.json", "sensibilidad.png")

# Model identifiers and their published expressions; no second model lives here.
BASE = {
    "sum_sf": "F1: sum_sf = suma(sf.values())",
    "E": "F2: E = B + 0.01 * sum_sf",
    "PM_ns": "F3: PM_ns = A * size**E * producto(em.values())",
    "PM": "F3: PM = PM_ns * sced_em",
    "phi": "F4: phi = D + 0.2 * (E - B)",
    "TDEV": "F5: TDEV = C * PM_ns**phi * sced_pct / 100",
    "t_h": "F6: t_h = c_pm / horas_por_pm",
    "C_trad": "F7: C_trad = PM * c_pm",
    "H_trad": "F8: H_trad = f * PM * horas_por_pm",
    "T": "F5 / T: T_meses indicado o TDEV",
}
ECONOMIA = {
    "C_adop": "F19: C_adop = C_fijo + u * H_trad * (1 - r)",
    "ahorro_bruto": "F20: ahorro_bruto = C_trad * f * r",
    "ahorro_neto": "F21: ahorro_neto = ahorro_bruto - C_adop",
    "rbc": "F22: rbc = ahorro_bruto / C_adop; None si C_adop = 0",
    "r_estrella": "F23: r_estrella = (C_fijo / H_trad + u) / (t_h + u)",
    "umbral_global": "F24: umbral_global = f * r_estrella",
    "c_adop_max": "F25: c_adop_max = ahorro_bruto",
    "margen": "F26: margen = r - r_estrella",
    "recuperacion_meses": "F27: T / rbc si rbc > 0; None en otro caso",
    "recupera_dentro_de_T": "F27: recuperacion_meses no nula y <= T",
    "C_ia": "F31: C_ia = C_trad - ahorro_neto",
}
ADOPCION = {
    "C_lic": "F17: C_lic = n_u * p_lic * T",
    "C_cap": "F17: C_cap = n_u * h_cap * t_h",
    "C_int": "F17: C_int = h_int * t_h",
    "C_fijo": "F17: C_fijo = C_lic + C_cap + C_int",
    "u": "F18: u = c_api_h + (rho_rev + rho_retr + delta * h_def) * t_h",
}
REAL = {
    "C_presup": "F32: suma(costo_usd de rubros en alcance)",
    "C_real": "F33: horas de programación * t_h + efectivo IA en alcance",
    "horas_totales": "F33: suma(horas_totales de sprints con programación)",
    "costo_ia_real": "F33: suma(licencias + api + capacitacion + integracion en alcance)",
    "H_real_4": "F9: suma(horas_4_actividades); None si falta alguna en alcance",
    "n_sprints": "F33: número de sprints con programación",
}
CONTRASTE = {
    "err_estimacion_ia": "F34: abs(C_real - C_ia) / C_real",
    "err_cocomo_vs_presup": "F35: abs(C_presup - C_trad) / C_presup",
    "desviacion_presup": "F36: (C_real - C_presup) / C_presup",
    "cumple_exactitud": "F34: err_estimacion_ia <= 0.25",
    "decision_estable": "F30 / F21: tres ahorros estrictamente del mismo signo; cero no estable",
    "verificacion_cocomo_ok": "F1–F5: verificación nominal independiente, tolerancia 0.1",
}
IA_FORMULAS = {
    "r": "F10",
    "rho_rev": "F13",
    "rho_retr": "F13",
    "c_api_h": "F14",
    "delta": "F15",
    "h_def": "F16",
    "p_lic": "F17",
    "h_cap": "F17",
    "h_int": "F17",
}


def _plain(value):
    if is_dataclass(value):
        return {
            f.name: _plain(getattr(value, f.name))
            for f in fields(value)
            if f.name != "ruta_interna"
        }
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, dict):
        return {k: _plain(v) for k, v in value.items()}
    if isinstance(value, (tuple, list)):
        return [_plain(v) for v in value]
    return value


class _Fuentes:
    def __init__(self, res):
        self.source = res.entradas.fuente
        self.cells = (
            {(c.hoja, c.celda): c.valor for c in self.source.celdas}
            if self.source
            else {}
        )
        self.p = {
            str(value): f"C{cell[1:]}"
            for (sheet, cell), value in self.cells.items()
            if sheet == "proyecto" and cell.startswith("B")
        }
        self.c = {
            str(value): f"B{cell[1:]}"
            for (sheet, cell), value in self.cells.items()
            if sheet == "constantes" and cell.startswith("A")
        }
        self.ia = {
            str(value): cell[1:]
            for (sheet, cell), value in self.cells.items()
            if sheet == "parametros_ia" and cell.startswith("A")
        }

    def locate(
        self, *, project=(), constants=(), parameters=(), case=False, runtime=()
    ):
        sources = []
        missing = []
        for symbol in project:
            cell = self.p.get(symbol)
            if cell and ("proyecto", cell) in self.cells:
                sources.append(f"proyecto!{cell}")
            else:
                missing.append(f"proyecto.{symbol}")
        for symbol in constants:
            if self.source and symbol in self.source.constantes_predeterminadas:
                sources.append(
                    f"predeterminado declarado: {symbol} (modelo/constantes.py)"
                )
            elif symbol in self.c:
                sources.append(f"constantes!{self.c[symbol]}")
            else:
                missing.append(f"constantes.{symbol}")
        for symbol in parameters:
            row = self.ia.get(symbol)
            if row:
                sources.append(f"parametros_ia!D{row}:I{row} ({symbol})")
            else:
                missing.append(f"parametros_ia.{symbol}")
        if case:
            locations = [
                f"{sheet}!{cell}"
                for (sheet, cell), value in self.cells.items()
                if sheet in ("presupuesto_inicial", "costo_real_sprint")
                and value not in (None, "")
            ]
            sources.extend(locations)
            if not locations:
                missing.append("caso")
        sources.extend(f"runtime: {name}" for name in runtime)
        if missing:
            sources.append("Fuente no disponible: " + ", ".join(missing))
        return (
            "; ".join(dict.fromkeys(sources))
            or "Fuente no disponible (constructor sin procedencia)"
        )


def _figure(value, formula, source):
    return {"valor": value, "formula": formula, "fuente": source}


def _annotate(contract, formulas, source):
    return {
        f.name: _figure(getattr(contract, f.name), formulas[f.name], source)
        for f in fields(contract)
    }


def _redact(value, private):
    """Source identities are basenames; private absolute paths never leave M8."""
    if isinstance(value, str):
        for path in private:
            if path:
                value = value.replace(path, "[ruta privada]").replace(
                    path.replace("\\", "/"), "[ruta privada]"
                )
        value = re.sub(
            r"(?<![\w])[A-Za-z]:[\\/][^\r\n;|\"<>]*", "[ruta privada]", value
        )
        value = re.sub(
            r"(?<![:\w])/(?:Users|home|tmp|var|mnt)/[^\r\n;|\"<>]*",
            "[ruta privada]",
            value,
        )
        return value
    if isinstance(value, dict):
        return {k: _redact(v, private) for k, v in value.items()}
    if isinstance(value, list):
        return [_redact(v, private) for v in value]
    return value


def a_diccionario(res: Resultados) -> dict:
    """Serialize full-precision snapshots, with honest transitive provenance."""
    src = _Fuentes(res)
    sf = tuple(res.entradas.proyecto.sf)
    em = tuple(res.entradas.proyecto.em)
    pm = ("Size", *sf, *em, "SCED_EM")
    core_constants = ("A", "B")
    horizon = ("T",) if res.entradas.proyecto.T_meses is not None else (*pm, "SCED_pct")
    horizon_constants = (
        () if res.entradas.proyecto.T_meses is not None else ("A", "B", "C", "D")
    )
    # A figure's input list includes upstream dependencies, not guessed coordinates.
    base_inputs = {
        "sum_sf": (sf, ()),
        "E": (sf, ("B",)),
        "PM_ns": (("Size", *sf, *em), core_constants),
        "PM": (pm, core_constants),
        "phi": (sf, ("D",)),
        "TDEV": ((*pm, "SCED_pct"), ("A", "B", "C", "D")),
        "t_h": (("c_pm",), ("horas_por_PM",)),
        "C_trad": ((*pm, "c_pm"), core_constants),
        "H_trad": ((*pm, "f"), (*core_constants, "horas_por_PM")),
        "T": (horizon, horizon_constants),
    }
    base = {
        name: _figure(
            getattr(res.base, name),
            formula,
            src.locate(project=base_inputs[name][0], constants=base_inputs[name][1]),
        )
        for name, formula in BASE.items()
    }
    parameters = tuple(IA_FORMULAS)
    all_project = (*pm, "SCED_pct", "c_pm", "f", "T", "n_u")
    all_constants = ("A", "B", "C", "D", "horas_por_PM")
    economic_source = src.locate(
        project=all_project, constants=all_constants, parameters=parameters
    )
    scenarios = []
    for scenario in res.decision.escenarios:
        scenarios.append(
            {
                "nombre": scenario.nombre,
                "valores": {
                    f.name: _figure(
                        getattr(scenario.valores, f.name),
                        f"{IA_FORMULAS[f.name]} / F28: selección {scenario.nombre}",
                        src.locate(parameters=(f.name,)),
                    )
                    for f in fields(scenario.valores)
                },
                "costos": _annotate(scenario.costos, ADOPCION, economic_source),
                "resultado": _annotate(scenario.resultado, ECONOMIA, economic_source),
            }
        )
    sens_source = src.locate(
        project=all_project,
        constants=all_constants,
        parameters=parameters,
        runtime=("x_pct",),
    )
    impacts = [
        {
            "parametro": row.parametro,
            "efecto_bajo": _figure(
                row.efecto_bajo, "F29: cambio ahorro neto al bajar x %", sens_source
            ),
            "efecto_alto": _figure(
                row.efecto_alto, "F29: cambio ahorro neto al subir x %", sens_source
            ),
        }
        for row in res.decision.sensibilidad
    ]
    real_source = src.locate(project=("c_pm",), constants=("horas_por_PM",), case=True)
    comparison_source = "; ".join(
        (
            economic_source,
            real_source,
            src.locate(
                constants=(
                    "PREC nominal",
                    "FLEX nominal",
                    "RESL nominal",
                    "TEAM nominal",
                    "PMAT nominal",
                    "EM nominal",
                    "SCED% nominal",
                )
            ),
            "runtime: verificacion_ok",
        )
    )
    expected = next((s for s in scenarios if s["nombre"] == "esperado"), None)
    summary = {
        "proyecto": res.entradas.proyecto.nombre,
        "archivo_entrada": res.entradas.fuente.nombre
        if res.entradas.fuente
        else "No disponible",
        "sha256_entrada": res.entradas.fuente.sha256 if res.entradas.fuente else None,
        "fecha_ejecucion": (res.fecha_ejecucion or date.today()).isoformat(),
        "naturaleza": "Supuestos declarados; no constituyen evidencia primaria. Respetar etiquetas sintéticas del proyecto.",
        "veredicto": _figure(
            res.decision.veredicto, "F30: regla pesimista/optimista", economic_source
        ),
        "C_trad": base["C_trad"],
        "mayor_efecto": _figure(
            impacts[0]["parametro"] if impacts else None,
            "F29: primer parámetro ordenado por mayor efecto absoluto",
            sens_source,
        ),
    }
    if expected:
        for name in (
            "ahorro_neto",
            "rbc",
            "r_estrella",
            "recuperacion_meses",
            "recupera_dentro_de_T",
        ):
            summary[name] = expected["resultado"][name]
        summary["r"] = expected["valores"]["r"]
    for scenario in scenarios:
        summary[f"C_ia_{scenario['nombre']}"] = scenario["resultado"]["C_ia"]
    summary["x_pct"] = _figure(
        res.decision.x_pct, "F29: porcentaje indicado", "runtime: x_pct"
    )
    data = {
        "resumen": summary,
        "base": base,
        "escenarios": scenarios,
        "sensibilidad": impacts,
        "costo_real": _annotate(res.costo_real, REAL, real_source)
        if res.costo_real
        else None,
        "contraste": _annotate(res.contraste, CONTRASTE, comparison_source)
        if res.contraste
        else None,
        "supuestos": list(res.entradas.supuestos),
        "avisos": [_plain(a) for a in res.entradas.avisos],
        "entradas": _plain(res.entradas),
    }
    private = ()
    if res.entradas.fuente:
        path = res.entradas.fuente.ruta_interna
        private = (path, str(Path(path).parent))
    data = _redact(data, private)
    # Strict JSON is also validation for non-finite input/results before I/O.
    json.dumps(data, ensure_ascii=False, allow_nan=False)
    return data


def _figures(value, prefix=""):
    if isinstance(value, dict):
        if set(value) == {"valor", "formula", "fuente"}:
            yield prefix, value
        else:
            for name, child in value.items():
                yield from _figures(child, f"{prefix}.{name}" if prefix else name)
    elif isinstance(value, list):
        for index, child in enumerate(value):
            yield from _figures(child, f"{prefix}.{index}")


def _format(name):
    tail = name.split(".")[-1]
    if tail in {"x_pct"}:
        return '0.0"%"'
    if tail.startswith("err_") or tail == "desviacion_presup":
        return "0.0%"
    if tail.startswith("C_") or tail in {
        "ahorro_bruto",
        "ahorro_neto",
        "c_adop_max",
        "efecto_bajo",
        "efecto_alto",
        "costo_ia_real",
    }:
        return '#,##0.00" USD"'
    return "0.0000"


def _display(name, value):
    if value is None:
        return "No disponible"
    if isinstance(value, bool):
        if name.endswith("decision_estable"):
            return "Estable" if value else "No estable"
        return "Sí" if value else "No"
    if isinstance(value, (int, float)):
        fmt = _format(name)
        if fmt == "0.0%":
            return f"{value:.1%}"
        if name.endswith("x_pct"):
            return f"{value:.1f}%"
        if "USD" in fmt:
            return f"{value:,.2f} USD"
        return f"{value:.4f}"
    return str(value)


def _literal(ws, values):
    """Explicit string data types prevent =,+,-,@ payload interpretation."""
    ws.append(values)
    for cell in ws[ws.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = "s"
        cell.font = Font(name="Arial", size=10)
        cell.alignment = Alignment(vertical="top", wrap_text=True)


def _workbook(data):
    w = Workbook()
    w.remove(w.active)
    names = ["resumen", "escenarios", "sensibilidad"]
    if data["contraste"] is not None:
        names.append("contraste")
    names += ["supuestos", "trazabilidad"]
    for name in names:
        ws = w.create_sheet(name)
        if name == "trazabilidad":
            _literal(ws, ["cifra", "fórmula", "entradas usadas", "valor"])
            rows = list(_figures(data))
        elif name == "supuestos":
            _literal(ws, ["tipo", "detalle", "fuente", "valor"])
            for assumption in data["supuestos"]:
                _literal(ws, ["Supuesto declarado", assumption, "parametros_ia", None])
            for warning in data["avisos"]:
                _literal(
                    ws,
                    [
                        warning["codigo"],
                        warning["mensaje"],
                        f"{warning['hoja']}!{warning['celda']}",
                        None,
                    ],
                )
            for cell in (
                data["entradas"].get("fuente", {}).get("celdas", [])
                if data["entradas"].get("fuente")
                else []
            ):
                _literal(
                    ws,
                    [
                        "Entrada conservada",
                        "Valor original en caché",
                        f"{cell['hoja']}!{cell['celda']}",
                        cell["valor"],
                    ],
                )
            if data["entradas"].get("fuente"):
                for constant in data["entradas"]["fuente"][
                    "constantes_predeterminadas"
                ]:
                    _literal(
                        ws,
                        [
                            "Predeterminado declarado",
                            constant,
                            "modelo/constantes.py",
                            None,
                        ],
                    )
            rows = []
        else:
            _literal(ws, ["cifra", "formula", "fuente", "valor"])
            if name == "resumen":
                for assumption in data["supuestos"]:
                    _literal(
                        ws,
                        [
                            "Supuesto declarado",
                            "Entrada declarada",
                            "parametros_ia",
                            assumption,
                        ],
                    )
                for warning in data["avisos"]:
                    _literal(
                        ws,
                        [
                            warning["codigo"],
                            "Aviso",
                            f"{warning['hoja']}!{warning['celda']}",
                            warning["mensaje"],
                        ],
                    )
                for key, value in data["resumen"].items():
                    if not isinstance(value, dict):
                        _literal(
                            ws,
                            [
                                key,
                                "Metadato de ejecución",
                                "runtime" if key == "fecha_ejecucion" else "entrada",
                                value,
                            ],
                        )
                rows = list(_figures(data["resumen"], "resumen")) + list(
                    _figures(data["base"], "base")
                )
            elif name == "contraste":
                rows = list(_figures(data["costo_real"], "costo_real")) + list(
                    _figures(data["contraste"], "contraste")
                )
            else:
                rows = list(_figures(data[name], name))
                labels = [
                    (
                        "parametro" if name == "sensibilidad" else "nombre",
                        row.get("parametro", row.get("nombre")),
                    )
                    for row in data[name]
                ]
                for index, (key, value) in enumerate(labels):
                    _literal(
                        ws,
                        [
                            f"{name}.{index}.{key}",
                            "F29" if name == "sensibilidad" else "F28",
                            "resultado del modelo",
                            value,
                        ],
                    )
        for key, figure in rows:
            _literal(ws, [key, figure["formula"], figure["fuente"], figure["valor"]])
            ws.cell(ws.max_row, 4).number_format = _format(key)
        for cell in ws[1]:
            cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="16324F")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col, width in zip("ABCD", (42, 65, 85, 30)):
            ws.column_dimensions[col].width = width
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
    buffer = BytesIO()
    try:
        w.save(buffer)
        return buffer.getvalue()
    finally:
        w.close()


def _md_escape(value):
    return str(value).replace("|", "\\|").replace("\n", "<br>").replace("\r", "")


def _markdown(data):
    lines = [
        "# Informe de estimación",
        "",
        "Instantánea estática; no es una hoja de cálculo ejecutable.",
        "",
    ]
    for key, value in data["resumen"].items():
        if not isinstance(value, dict):
            lines.append(f"- **{key}**: {_md_escape(value)}")
    for title, keys in (
        ("Resumen", ("resumen", "base")),
        ("Escenarios", ("escenarios",)),
        ("Sensibilidad", ("sensibilidad",)),
        ("Contraste", ("costo_real", "contraste")),
    ):
        if title == "Contraste" and data["contraste"] is None:
            continue
        if title in ("Escenarios", "Sensibilidad"):
            section = keys[0]
            label = "nombre" if title == "Escenarios" else "parametro"
            lines += [
                "",
                f"{title}: "
                + "; ".join(
                    f"{i}: {_md_escape(row[label])}"
                    for i, row in enumerate(data[section])
                ),
            ]
        lines += [
            "",
            f"## {title}",
            "",
            "| cifra | valor | fórmula | fuente |",
            "|---|---:|---|---|",
        ]
        for section in keys:
            for name, figure in _figures(data[section], section):
                lines.append(
                    "| "
                    + " | ".join(
                        _md_escape(v)
                        for v in (
                            name,
                            _display(name, figure["valor"]),
                            figure["formula"],
                            figure["fuente"],
                        )
                    )
                    + " |"
                )
    lines += ["", "## Supuestos", ""]
    lines += [f"- {_md_escape(s)}" for s in data["supuestos"]] or ["- Ninguno"]
    lines += ["", "## Avisos", ""]
    lines += [
        f"- [{a['codigo']}] {_md_escape(a['hoja'])}!{a['celda']}: {_md_escape(a['mensaje'])}"
        for a in data["avisos"]
    ] or ["- Ninguno"]
    lines += [
        "",
        "## Entradas y fuentes conservadas",
        "",
        "```json",
        json.dumps(data["entradas"], ensure_ascii=False, allow_nan=False, indent=2),
        "```",
        "",
    ]
    return "\n".join(lines).encode("utf-8")


def _png(data):
    rows = data["sensibilidad"]
    fig = Figure(figsize=(11, max(5, len(rows) * 0.43)), dpi=120, layout="constrained")
    canvas = FigureCanvasAgg(fig)
    axes = fig.subplots()
    positions = list(range(len(rows)))
    axes.barh(
        [y - 0.18 for y in positions],
        [r["efecto_bajo"]["valor"] for r in rows],
        height=0.34,
        color="#2D6A8A",
        label="Variación baja",
    )
    axes.barh(
        [y + 0.18 for y in positions],
        [r["efecto_alto"]["valor"] for r in rows],
        height=0.34,
        color="#CA7B36",
        label="Variación alta",
    )
    axes.set_yticks(positions, [r["parametro"] for r in rows], fontname="Arial")
    axes.invert_yaxis()
    axes.axvline(0, color="#444444", linewidth=0.8)
    axes.set_xlabel("Cambio del ahorro neto (USD)", fontname="Arial")
    axes.set_title(
        f"Sensibilidad ±{data['resumen']['x_pct']['valor']:.1f}% — {data['resumen']['proyecto']}",
        fontname="Arial",
    )
    axes.legend(prop={"family": "Arial"})
    axes.grid(axis="x", alpha=0.2)
    for tick in axes.get_xticklabels():
        tick.set_fontname("Arial")
    output = BytesIO()
    canvas.print_png(output)
    return output.getvalue()


def escribir_reporte(res: Resultados, carpeta: Path) -> tuple[Path, ...]:
    """Preflight all targets, stage in memory, then exclusively create four files.

    Failed writes remove only files successfully created by this call. Existing
    output bytes, unrelated siblings and the protected input are never replaced.
    """
    carpeta = Path(carpeta)
    paths = tuple(carpeta / name for name in ARCHIVOS)
    source = (
        Path(res.entradas.fuente.ruta_interna).resolve()
        if res.entradas.fuente
        else None
    )
    for path in paths:
        if (
            path.exists()
            or path.is_symlink()
            or (source is not None and path.resolve() == source)
        ):
            raise FileExistsError(f"Destino existente o protegido: {path.name}")
    data = a_diccionario(res)
    payloads = (
        _workbook(data),
        _markdown(data),
        json.dumps(data, ensure_ascii=False, allow_nan=False, indent=2).encode("utf-8"),
        _png(data),
    )
    carpeta.mkdir(parents=True, exist_ok=True)
    owned = []
    try:
        for path, payload in zip(paths, payloads):
            with path.open("xb") as stream:
                owned.append(path)
                stream.write(payload)
    except BaseException:
        for path in reversed(owned):
            path.unlink()
        raise
    return paths
